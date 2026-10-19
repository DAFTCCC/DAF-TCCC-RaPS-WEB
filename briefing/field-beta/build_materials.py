#!/usr/bin/env python3
"""Build the leadership deck and printable instructor beta materials.

Run from this directory: python3 build_materials.py
Requires python-pptx, ReportLab, Pillow and python-docx. Convert the presentation
to PDF with LibreOffice after building; the README documents the exact command.
No application, database, account or network operation is performed.
"""
from pathlib import Path
import csv
import json
import re
from xml.sax.saxutils import escape

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
APP_URL = 'https://daftccc.github.io/DAF-TCCC-RaPS-WEB/'
SOURCE_COMMIT = '104aaae9031cbe4c361022ebd380ff169bfca2c9'
SOURCE_URL = 'https://github.com/DAFTCCC/DAF-TCCC-RaPS-WEB/blob/' + SOURCE_COMMIT + '/'
PATCH_SOURCE_URL = 'https://github.com/DAFTCCC/DAF-TCCC-RaPS-WEB/blob/codex/field-beta-briefing/'
NAVY = '102A43'
TEAL = '007F86'
ORANGE = 'E99A3C'
PALE = 'F3F6FA'
INK = '17354B'
MUTED = '516677'
WHITE = 'FFFFFF'
RED = 'B64040'
LINE = 'D9E2EB'
PPTX_PATH = HERE / 'RaPS-Leadership-Beta-Brief.pptx'
FONT_PATH = '/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'
if not Path(FONT_PATH).is_file():
    FONT_PATH = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, line=None, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def estimate_lines(text, size, width):
    font = ImageFont.truetype(FONT_PATH, round(size * 96 / 72))
    limit = max(1, (width * 96) - 14)
    count = 0
    for source in text.split('\n'):
        current = ''
        for word in source.split():
            if current and font.getlength(current + ' ' + word) > limit:
                count += 1
                current = word
            else:
                current = (current + ' ' + word).strip()
        count += 1
    return count


def text(slide, x, y, w, h, value, size=22, color=INK, bold=False, minimum=17):
    while size > minimum and estimate_lines(value, size, w) * size * 1.18 > (h * 72 - 8):
        size -= 1
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(2)
    frame.margin_top = frame.margin_bottom = Pt(2)
    for index, line in enumerate(value.split('\n')):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = line
        para.font.name = 'Arial'
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = rgb(color)
        para.space_after = Pt(5)
        para.line_spacing = 1.08
    return shape


def notes(slide, message, refs='DEVELOPMENT.md; PERMISSIONS.md; WEB8-DEPLOYMENT.md'):
    slide.notes_slide.notes_text_frame.text = message + '\n\nSource basis: ' + refs + '\n' + SOURCE_URL


def new_slide(prs, title, section, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, NAVY if dark else PALE)
    rect(slide, 0, 0, 0.13, 7.5, TEAL)
    text(slide, .55, .32, 12.1, .32, section.upper(), 11, ORANGE if dark else TEAL, True, 11)
    text(slide, .55, .83, 12.1, 1.05, title, 32, WHITE if dark else NAVY, True, 27)
    rect(slide, .57, 6.99, 12.16, .013, '355369' if dark else LINE)
    text(slide, .57, 7.10, 11.3, .23, 'RaPS  |  Instructor break-it beta  |  Fictitious beta records  |  web.9',
         9, 'B9CBD8' if dark else MUTED, minimum=9)
    text(slide, 12.12, 7.08, .5, .25, f'{len(prs.slides):02d}', 10,
         'B9CBD8' if dark else MUTED, minimum=10)
    return slide


def card(slide, x, y, w, h, label, body, accent=TEAL, body_size=21):
    rect(slide, x, y, w, h, WHITE, LINE, True)
    rect(slide, x+.18, y+.22, .065, h-.44, accent)
    text(slide, x+.39, y+.22, w-.65, .68, label, 23, NAVY, True, 20)
    text(slide, x+.39, y+.97, w-.65, h-1.12, body, body_size, MUTED, minimum=18)


def band(slide, value, y=6.28, fill=TEAL, size=20):
    rect(slide, .57, y, 12.16, .54, fill, rounded=True)
    text(slide, .78, y+.07, 11.75, .41, value, size, WHITE, True, 17)


def build_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = 'RaPS Instructor Break-It Beta — Leadership Brief'
    prs.core_properties.subject = 'Controlled TEST pilot, instructor challenges, feedback and readiness decisions'
    prs.core_properties.author = 'RaPS pilot planning'
    prs.core_properties.keywords = 'RaPS, TEST, beta, instructors, MAJCOM, field evaluation'

    slide = new_slide(prs, '', 'Leadership and MAJCOM launch brief', True)
    text(slide, .65, 1.55, 8.45, 2.15, 'Let instructors\ntry to break RaPS.', 46, WHITE, True, 42)
    text(slide, .68, 4.02, 8.25, 1.0, 'A supervised field beta focused on failures, friction and useful feedback.', 25, 'C7D8E4', minimum=24)
    slide.shapes.add_picture(str(ROOT / 'assets/app-icon.png'), Inches(9.52), Inches(2.02), width=Inches(2.85))
    band(slide, 'Decision requested: authorize a small TEST pilot', 5.88, TEAL, 23)
    notes(slide, 'Open with the challenge: we want instructors to use their normal habits, interrupt the workflow, make controlled mistakes and tell us what gets in the way. This is a proposed supervised TEST pilot. It is not a declaration that field reliability or the assessment instrument has already been validated. Allow about ten minutes for the brief, then protect forty minutes for instructor challenges.')

    slide = new_slide(prs, 'Clear the start blocker before the pilot', 'Readiness recommendation')
    card(slide, .57, 2.03, 3.94, 3.74, 'Engineering checks', 'Startup failure reproduced.\nRegression checks cover the shared fix and normal TEST behavior.', TEAL, 21)
    card(slide, 4.69, 2.03, 3.94, 3.74, 'Deployed foundation', 'Web.8 backend verified.\nWeb.9 frontend correction needs publication and a deployed retest.', TEAL, 21)
    card(slide, 8.81, 2.03, 3.94, 3.74, 'Field evidence needed', 'Retry the reported class.\nReal devices and recovery.\nTwo-user results and retention.', ORANGE, 21)
    band(slide, 'Hold the cohort until web.9 is published and the owner dry run passes.')
    notes(slide, 'The owner supplied a successful web.8 NUC backend rollout. An actual assessment-start failure was then reported: INSERT RETURNING could not satisfy the class read helper within the insertion snapshot. The same failure was reproduced for ordinary and TEST classes. Web.9 inserts and separately verifies the readable revision without changing RLS. It also corrects TEST to the ordinary supported-command and physical-installation workflow. Local regression evidence does not prove the reported NUC class has been retried. Publish the frontend, confirm web.9 on each device, and complete the owner dry run before inviting the cohort. No NUC migration or Auth restart is required for this patch.')

    slide = new_slide(prs, 'What each command provides', 'Pilot proposal — adjust locally')
    card(slide, .57, 1.94, 5.98, 1.88, '1 PM + 2 instructors', 'A site lead, hands-on testers and an observer when practical.', TEAL, 20)
    card(slide, 6.75, 1.94, 5.98, 1.88, 'Their normal devices', 'Phone/tablet, browser/PWA, old and new hardware, ordinary connectivity.', TEAL, 20)
    card(slide, .57, 4.02, 5.98, 1.88, 'Protected testing time', 'A 60-minute launch, then follow-up challenges during the week.', ORANGE, 20)
    card(slide, 6.75, 4.02, 5.98, 1.88, 'A feedback owner', 'Instructor → site PM → pilot owner; daily triage and a retest slot.', ORANGE, 20)
    band(slide, 'Start with 2–3 sites over one week. Expand the cohort after review.', 6.30)
    notes(slide, 'This is a suggested cohort and duration, not an approved tasking or a statistically powered validation study. Include both experienced and newer instructors and a practical device mix. Name an Enterprise/NUC recovery owner and an SME. PMs collect findings through the supplied form and issue log; there is no newly created online feedback endpoint. Temporary credentials must be delivered privately. Do not share accounts.')

    slide = new_slide(prs, 'TEST is one more supported MAJCOM', 'Use the ordinary environment and workflow')
    card(slide, .57, 2.02, 5.98, 2.58, 'Same installations', 'Select TEST, then an existing installation.\nAll 73 physical locations are available; no new duplicate locations.', TEAL, 22)
    card(slide, 6.75, 2.02, 5.98, 2.58, 'Same training workflow', 'Create, start, observe, sync, finalize and close through the normal controls and permission rules.', ORANGE, 22)
    text(slide, .75, 4.87, 11.9, .70, 'All MAJCOMs includes TEST. Use fictitious learners and clearly named beta classes.', 23, NAVY, True)
    text(slide, .75, 5.81, 11.9, .77, 'Supported MAJCOM is a training label. Assigned installation roles and explicit sharing control access.', 20, MUTED, minimum=19)
    notes(slide, 'The prior release added duplicate TEST-* bases, which exceeded the requested catalog change. Web.9 stops offering those for new classes or access grants, but preserves their saved records and memberships. New TEST classes use existing physical installation IDs. TEST is not a separate server or permission role, and selecting it grants no access. A TEST MAJCOM-manager membership does not automatically grant oversight over other host commands or permission to grade. Use the normal assigned evaluator/PM memberships and ordinary reporting scope. Current dashboards group cloud reports by the base parent command, as for other supported-command labels. Existing duplicate-base accounts are not silently moved to a physical base; an administrator must review any new grant.')

    slide = new_slide(prs, 'Owner dry run before instructors arrive', 'Four preflight gates')
    rows = [
        ('01', 'Release on the device', 'Every participating browser/PWA shows 3.4.12-web.9.'),
        ('02', 'Real sign-in and recovery', 'Temporary password → new password → fresh login; reset and scope checks.'),
        ('03', 'A known case on two devices', 'Claim, save, sync, finalize and reconcile the stored result with the SME case.'),
        ('04', 'Support and retention ownership', 'Named support path, official-process fallback, both archive workflows checked.'),
    ]
    for i, (number, label, body) in enumerate(rows):
        y = 1.96+i*1.04
        rect(slide, .61, y+.04, .59, .59, TEAL, rounded=True)
        text(slide, .67, y+.12, .48, .4, number, 19, WHITE, True)
        text(slide, 1.42, y-.01, 10.95, .40, label, 22, NAVY, True)
        text(slide, 1.42, y+.43, 10.95, .49, body, 19, MUTED, minimum=18)
    text(slide, .66, 6.30, 12.05, .48, 'Before meaningful record retention or expansion: verify off-NUC backup and an isolated restore.', 18, MUTED, minimum=17)
    notes(slide, 'Record evidence rather than checking these boxes by assumption. The successful backend health/catalog/401 smoke did not create an account or complete a training case. Retain a second established administrator. If email accounts participate, verify their actual recovery delivery separately. Two archive paths exist: the pre-existing automatic closed-class retention worker and the new selected-class reporting archive worker. Verify both separately. Rehearse restoration into an isolated target; never restore over production to test it.')

    slide = new_slide(prs, 'Give instructors permission to be difficult', 'The break-it challenge')
    card(slide, .57, 1.96, 5.98, 1.97, 'PRESS THE INTERFACE', 'Mistap. Undo. Enter long notes. Rotate. Use gloves and the keyboard.', TEAL, 20)
    card(slide, 6.75, 1.96, 5.98, 1.97, 'INTERRUPT THE SESSION', 'Background. Lock. Return. Reload. Check active timer recovery.', ORANGE, 20)
    card(slide, .57, 4.12, 5.98, 1.97, 'LOSE THE CONNECTION', 'Claim online, disconnect, record, reconnect and compare exact entries.', TEAL, 20)
    card(slide, 6.75, 4.12, 5.98, 1.97, 'COLLIDE WORKFLOWS', 'Two devices. Two evaluators. Stale views. Final records and scoped reports.', ORANGE, 20)
    band(slide, 'Use ordinary interface actions and disposable TEST cases. Capture what happens.', 6.32)
    notes(slide, 'The instructor card contains eight missions. Assign four per tester during the forty-minute challenge window, including five minutes of free exploration. Cover all eight across the cohort and week. Do not force full finalization or closure into the launch: the case must satisfy existing standards and required timing. Maintain one active assessment claim per evaluator; reuse the same claimed case for observation, outage and reload missions. Use separate scheduled fixtures for competing claims and finalization. Do not erase unsynced browser storage, alter clinical standards or stress the infrastructure.')

    slide = new_slide(prs, 'Expected safeguards are part of the test', 'Brief these before calling something a bug')
    items = [
        ('Offline work', 'Continue a previously claimed assessment. New starts, finalization, void and close require the server.'),
        ('Actual timer reload', 'Recovery requires VOID with a reason, then a new instance and a repeated approved timed performance.'),
        ('Finalized results', 'Finalized assessment fields stay read-only. Coaching does not revise an official result.'),
        ('Local copy removal', 'Removing one browser copy keeps the shared server record. Retention controls are owner-operated.'),
    ]
    for i, (label, body) in enumerate(items):
        y = 1.97+i*1.11
        rect(slide, .61, y, 12.08, 1.00, WHITE, LINE, True)
        text(slide, .82, y+.17, 2.48, .64, label, 21, TEAL, True)
        text(slide, 3.43, y+.15, 8.98, .76, body, 20, INK, minimum=19)
    notes(slide, 'A timer that survives a foreground update is different from one recovered after an actual browser reload or process termination. The latter cannot be resumed or stopped as a trusted uninterrupted clock; the application requires VOID/restart. A void reason and new instance must preserve audit history. Do not claim the device always keeps timing through a crash. Cached offline authority must still be valid. Competing claim denial and locked final results can be successful tests, not defects. Unexpected data loss, falsely trusted elapsed time or false finalization are stop conditions. Timer sources: app.js timerNeedsRecovery/timerAction/timerMarkup/renderActiveTimers/updateLiveClocks.', 'app.js:1846–1942; PERMISSIONS.md; DEVELOPMENT.md')

    slide = new_slide(prs, 'Turn “it broke” into a fixable report', 'One report per finding')
    labels = ['TRY', 'CAPTURE', 'COMPARE', 'RETURN']
    for i, label in enumerate(labels):
        x = .65+i*3.08
        rect(slide, x, 2.02, 2.76, .68, TEAL if i != 2 else ORANGE, rounded=True)
        text(slide, x+.13, 2.19, 2.5, .39, label, 22, WHITE, True)
    text(slide, .77, 2.97, 11.86, 1.0, 'Record the device, version, role/site, TEST case, exact steps, expected versus actual result, and local/cloud state.', 25, NAVY, True, 23)
    rect(slide, .64, 4.22, 12.04, 1.42, WHITE, LINE, True)
    text(slide, .84, 4.38, 11.64, .30, 'ILLUSTRATIVE REPORT — NOT AN OBSERVED DEFECT', 12, TEAL, True, 12)
    text(slide, .84, 4.83, 11.64, .61, '“After rotation, STOP did not respond. Timer remained active. Reproduced 2/3 tries on [device/browser].”', 23, INK, minimum=21)
    band(slide, 'Return TEST-only evidence to the site PM. Include one thing worth keeping.', 6.28)
    notes(slide, 'Use the fillable feedback PDF for each finding and the CSV issue log for site consolidation. Capture the local time and zone to help correlate events. Record independent/helped/blocked/not-run rather than only a satisfaction score. A screenshot of synthetic data is useful; passwords, tokens, recovery links and real learner data are not. The sample issue is invented to illustrate the level of detail, not a known defect. There is no new in-app feedback service: the site PM routes reports to the named pilot owner.')

    slide = new_slide(prs, 'Pause the affected lane when evidence is at risk', 'Triage without losing the useful evidence')
    card(slide, .57, 2.02, 3.94, 3.90, 'PAUSE + ESCALATE', 'Unexpected access.\nWrong or lost grading.\nSilent overwrites.\nFalse save/final score.\nLive records affected.', RED, 21)
    card(slide, 4.69, 2.02, 3.94, 3.90, 'CAPTURE + TRIAGE', 'Blocked workflow.\nUnclear recovery.\nUnreachable control.\nMisleading status.\nRepeated workaround.', ORANGE, 21)
    card(slide, 8.81, 2.02, 3.94, 3.90, 'RECORD + IMPROVE', 'Friction and preferences.\nMissing capability.\nUseful display choice.\nGood recovery.\nWhat should stay.', TEAL, 21)
    band(slide, 'Keep the device state. Site PM → pilot owner → Enterprise/NUC support.', 6.30)
    notes(slide, 'Pause only the affected lane while the owner assesses a critical observation, unless the issue is systemic. Preserve state and recover/export with support before clearing browser storage or reinstalling. Continue the official training assessment through the established process. Do not repeatedly retry a destructive action or choose LOCAL/SERVER without comparing the branches. Reopen a paused mission only after owner review. A disagreement with an SME reference case also needs review; the beta must not quietly redefine a clinical rule.')

    slide = new_slide(prs, 'Decide on evidence, then retest the fixes', 'Proposed local expansion gates')
    gates = [
        ('No open critical blocker', 'Access, evidence integrity and stored final results behave correctly.'),
        ('Known cases reconcile', 'Two devices, exports and scoped TEST reports agree on the selected fixtures.'),
        ('Recovery is demonstrated', 'Real account recovery, outage/reconnect and timer interruption recovery work.'),
        ('Instructors can work', 'Report independent/helped/blocked counts, observation burden and concrete friction.'),
    ]
    for i, (label, body) in enumerate(gates):
        y = 1.98+i*1.02
        text(slide, .74, y, 11.9, .37, label, 23, TEAL, True)
        text(slide, .74, y+.42, 11.9, .49, body, 20, MUTED, minimum=19)
    band(slide, 'Proposed operational gates. Clinical agreement needs SME review of the same cases.', 6.30)
    notes(slide, 'Use counts and denominators for this small formative cohort. Do not represent a satisfaction score, pass-rate difference or a small pilot as clinical validation. Compare instructor judgments on the same approved cases; investigate critical-error disagreement separately from contributor hypotheses. Compare paper/app order across equivalent lanes when practical. Expansion should wait for recovery evidence and critical fixes. Leadership may extend a TEST pilot for usability gaps with owners and retest dates.')

    slide = new_slide(prs, 'A short cycle with protected instructor time', 'Proposed one-week pilot')
    stages = [
        ('DAY 0', 'Owner dry run', 'Accounts, scopes, known case, support and recovery evidence.'),
        ('DAY 1', 'Brief + challenge', '10-minute brief, 40-minute break-it session, 10-minute debrief.'),
        ('DAYS 2–5', 'Use + report', 'Cover remaining missions; use normal devices; daily PM triage.'),
        ('DAYS 6–7', 'Fix + retest', 'Retest affected devices and cases; brief hold, extend or expand.'),
    ]
    for i, (day, label, body) in enumerate(stages):
        x = .59+i*3.08
        text(slide, x+.08, 2.03, 2.72, .44, day, 19, TEAL, True)
        card(slide, x, 2.67, 2.92, 3.20, label, body, ORANGE if i == 3 else TEAL, 21)
    notes(slide, 'This relative schedule is a proposal. Assign actual dates and named people before the brief. The sixty-minute launch does not require every tester to complete all eight missions or all clinical requirements. Use follow-up slots for finalization/closure and the owner-only retention demonstration. Run one short daily triage and return fixes to the original reporter for retest. Give instructors time away from the observation task to report findings rather than adding the burden invisibly during a real lane.')

    slide = new_slide(prs, 'The leadership ask', 'Authorize the pilot and name the owners', True)
    text(slide, .72, 2.01, 11.82, 2.68, 'Approve 2–3 TEST sites and a device mix.\nNominate each PM and two instructors.\nName the pilot owner, Enterprise support and SME.\nProtect time to test, report and retest.', 28, WHITE, True, 26)
    text(slide, .75, 5.05, 11.76, .8, 'Return: a ranked defect list, instructor recommendations and a scoped go / hold / expand decision.', 24, 'C7D8E4', minimum=23)
    band(slide, 'Start small. Try to break it. Bring evidence.', 6.24, TEAL, 24)
    notes(slide, 'Close by asking leaders to resource a supervised pilot, not to endorse a finished product. Confirm the named support route and when the site PM will return the consolidated log. The intended deliverable is evidence and a prioritized improvement list, including what instructors want to keep. Current official documentation/process remains authoritative during the beta. App: ' + APP_URL)

    slide = new_slide(prs, 'Appendix: who does what', 'Roles for the beta')
    role_rows = [
        ('Instructor / evaluator', 'Work at the assigned installation; hold the assessment claim; challenge and report.'),
        ('Installation PM', 'Prepare fixtures/assignments, pair testers, reconcile counts and collect findings.'),
        ('MAJCOM Manager', 'Review reports within ordinary authorized command scope and track site feedback.'),
        ('Enterprise / NUC owner', 'Provision/reset accounts; operate retention; preserve backups and recovery evidence.'),
        ('SME / pilot owner', 'Set reference cases, review rule disagreements, triage and authorize retesting.'),
    ]
    for i, (label, body) in enumerate(role_rows):
        y = 1.97+i*.92
        text(slide, .73, y, 3.25, .64, label, 20, TEAL, True, 19)
        text(slide, 4.10, y, 8.42, .70, body, 20, INK, minimum=19)
    notes(slide, 'A named person can have more than one authorized membership, but a role label alone does not grant the other memberships. MAJCOM oversight follows ordinary assigned scope; selecting TEST grants no new access, grading or account administration. Hands-on participants need the operational membership and any necessary assignment. Existing explicit cross-scope assignments retain their ordinary sharing semantics; use only the facilitator-approved fixtures.')

    slide = new_slide(prs, 'Appendix: remaining operational checks', 'Items to prove before expansion')
    checks = [
        ('Real device release + identity', 'Browser/PWA version, first login, reset and scope on actual devices.'),
        ('Outages + ownership', 'Local persistence, claim exclusion, reconnect and exact stored results.'),
        ('Physical usability + timing', 'Gloves, light, keyboard, interruption recovery and reference-clock comparison.'),
        ('Both retention workflows', 'Existing automatic closed-class worker; selected-class reporting archive worker.'),
        ('Recoverability + support', 'Off-NUC backup, isolated restore, owner escalation and field support process.'),
    ]
    for i, (label, body) in enumerate(checks):
        y = 1.97+i*.90
        text(slide, .73, y, 4.00, .62, label, 20, NAVY, True, 19)
        text(slide, 4.86, y, 7.67, .70, body, 19, MUTED, minimum=18)
    notes(slide, 'The reported startup defect has a local fix; deployed retesting is still required. The following items are additional evidence gaps and operational dependencies. Detailed gates and source references are in facilitator-guide.pdf and README.md. Clinical wording and scoring/lifecycle algorithms were protected by regression comparisons. Physical device use and same-case evaluator agreement still require the pilot. Two retention workers must not be conflated: closed-class automatic retention is pre-existing, while a selected reporting archive needs an Enterprise request and owner-run worker verification. Sources: DEVELOPMENT.md, PERMISSIONS.md, WEB8-DEPLOYMENT.md, app.js. Main source basis: ' + SOURCE_COMMIT)

    prs.save(PPTX_PATH)
    return len(prs.slides)


def register_pdf_fonts():
    base = Path('/usr/share/fonts/truetype/dejavu')
    pdfmetrics.registerFont(TTFont('BetaSans', str(base / 'DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('BetaSans-Bold', str(base / 'DejaVuSans-Bold.ttf')))
    pdfmetrics.registerFont(TTFont('BetaSans-Oblique', str(base / 'DejaVuSans-Oblique.ttf')))
    pdfmetrics.registerFontFamily('BetaSans', normal='BetaSans', bold='BetaSans-Bold', italic='BetaSans-Oblique')


def inline(source):
    value = escape(source)
    value = re.sub(r'`([^`]+)`', r'<font color="#007F86">\1</font>', value)
    value = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', value)
    def link(match):
        label, target = match.group(1), match.group(2)
        if not target.startswith('http'):
            path = target.replace('../../', '')
            target = (PATCH_SOURCE_URL if path == 'WEB9-DEPLOYMENT.md' else SOURCE_URL) + path
        return '<link href="' + target + '" color="#007F86">' + label + '</link>'
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)', link, value)


def build_markdown_pdf(source, target, label):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('BetaBody', fontName='BetaSans', fontSize=10.0, leading=13.4,
                              textColor=colors.HexColor('#17354B'), spaceAfter=4))
    styles.add(ParagraphStyle('BetaH1', parent=styles['BetaBody'], fontName='BetaSans-Bold',
                              fontSize=19, leading=23, spaceAfter=8))
    styles.add(ParagraphStyle('BetaH2', parent=styles['BetaBody'], fontName='BetaSans-Bold',
                              fontSize=11.4, leading=14.1, spaceBefore=5, spaceAfter=4,
                              textColor=colors.HexColor('#007F86'), keepWithNext=True))
    story, pending = [], []
    def flush():
        if pending:
            story.append(Paragraph(inline(' '.join(pending)), styles['BetaBody']))
            pending.clear()
    for raw in source.read_text().splitlines():
        line = raw.strip()
        if line == '<!-- pagebreak -->':
            flush(); story.append(PageBreak())
        elif line.startswith('# '):
            flush(); story.append(Paragraph(inline(line[2:]), styles['BetaH1']))
        elif line.startswith('## '):
            flush(); story.append(Paragraph(inline(line[3:]), styles['BetaH2']))
        elif not line:
            flush()
        elif line.startswith('- ') or re.match(r'^\d+\. ', line):
            flush(); story.append(Paragraph(inline(line), styles['BetaBody']))
        else:
            pending.append(line)
    flush()
    def page(c, doc):
        width, height = letter
        c.setFillColor(colors.HexColor('#007F86'))
        c.rect(36, height-32, width-72, 3, fill=1, stroke=0)
        c.setFont('BetaSans', 8)
        c.setFillColor(colors.HexColor('#516677'))
        c.drawString(36, height-23, 'RaPS | INSTRUCTOR BREAK-IT BETA | ' + label.upper())
        c.drawString(36, 24, 'Fictitious beta records • web.9 • Return feedback to your nominated site PM')
        c.drawRightString(width-36, 24, str(doc.page))
    doc = SimpleDocTemplate(str(target), pagesize=letter, rightMargin=36, leftMargin=36,
                            topMargin=46, bottomMargin=41, title=label, author='RaPS pilot planning')
    doc.build(story, onFirstPage=page, onLaterPages=page)


def build_feedback():
    path = HERE / 'RaPS-Beta-Feedback-Form.pdf'
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setTitle('RaPS instructor beta — one finding per form')
    c.setAuthor('RaPS pilot planning')
    c.setFillColor(colors.HexColor('#102A43'))
    c.rect(0, 695, 612, 97, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont('BetaSans-Bold', 21); c.drawString(36, 753, 'Tell us what happened.')
    c.setFont('BetaSans', 10); c.drawString(36, 730, 'One finding per form • TEST records only • Include what you want to keep')
    c.setFont('BetaSans', 9); c.drawString(36, 711, 'Fill in a PDF viewer or print. Return to your nominated site Program Manager.')
    fields = []
    def field(name, label, x, y, w, h=23, multiline=False):
        c.setFillColor(colors.HexColor('#17354B')); c.setFont('BetaSans-Bold', 8.4)
        c.drawString(x, y+h+5, label)
        c.acroForm.textfield(name=name, tooltip=label, x=x, y=y, width=w, height=h,
                            fontName='Helvetica', fontSize=10, borderWidth=.7,
                            borderColor=colors.HexColor('#B8C8D4'), fillColor=colors.HexColor('#F3F6FA'),
                            textColor=colors.HexColor('#17354B'), forceBorder=True,
                            fieldFlags='multiline' if multiline else '', maxlen=2000 if multiline else 180)
        fields.append(name)
    field('report_id', 'Report ID / tester alias', 36, 648, 258)
    field('when', 'Local time and time zone', 312, 648, 264)
    field('site_role', 'TEST installation and role', 36, 598, 258)
    field('version', 'App version / browser or installed PWA', 312, 598, 264)
    field('device', 'Device model / OS / browser version', 36, 548, 258)
    field('network', 'Connectivity before / during / after', 312, 548, 264)
    field('mission_case', 'Mission number / synthetic class or case ID', 36, 498, 540)
    c.setFont('BetaSans', 9.3); c.setFillColor(colors.HexColor('#17354B'))
    c.drawString(36, 474, 'Impact:')
    for name, label, x in [('blocked','Blocked',82),('workaround','Workaround',172),
                           ('friction','Friction',291),('positive','What works',386),('idea','Idea',506)]:
        c.acroForm.checkbox(name=name, tooltip=label, x=x, y=470, size=11,
                            borderWidth=.8, borderColor=colors.HexColor('#516677'),
                            fillColor=colors.white, buttonStyle='check')
        c.drawString(x+16, 473, label)
        fields.append(name)
    field('steps', 'Exact steps to reproduce, in order', 36, 359, 540, 86, True)
    field('expected', 'What did you expect?', 36, 280, 258, 54, True)
    field('actual', 'What actually happened?', 312, 280, 264, 54, True)
    field('state', 'Local/cloud status; missing or retained evidence; recovery tried', 36, 202, 540, 53, True)
    field('repro_help', 'Repeated how many times? Independent / help / blocked / not run?', 36, 150, 540)
    field('suggestion', 'Suggested change / one thing to keep / screenshot reference', 36, 69, 540, 53, True)
    c.setFont('BetaSans', 8.3); c.setFillColor(colors.HexColor('#516677'))
    c.drawString(36, 42, 'No passwords, tokens, recovery links or real learner details. Preserve unsynced device state.')
    c.drawString(36, 27, 'Pause affected testing for unexpected access, lost/wrong grading or false save/finalization success.')
    c.showPage(); c.save()
    (HERE / 'feedback-fields.json').write_text(json.dumps(fields, indent=2)+'\n')
    document = Document()
    document.sections[0].top_margin = DocInches(.6)
    document.sections[0].bottom_margin = DocInches(.6)
    document.add_heading('RaPS instructor beta feedback', 0)
    document.add_paragraph('One finding per form. TEST data only. Return to your nominated site PM.')
    prompts = ['Report ID / tester alias; local time and zone', 'TEST site and role; mission / synthetic case ID',
               'Device / OS / browser / PWA; app version; network state',
               'Impact: blocked / workaround / friction / what works / idea',
               'Exact steps to reproduce', 'Expected result', 'Actual result',
               'Local/cloud state, evidence retained or missing, recovery tried',
               'Repeatability; independent / helped / blocked / not run',
               'Suggested change; one thing to keep; sanitized screenshot reference']
    table = document.add_table(rows=0, cols=1)
    table.style = 'Table Grid'
    for prompt in prompts:
        cell = table.add_row().cells[0]
        cell.paragraphs[0].add_run(prompt).bold = True
        cell.add_paragraph(' ')
    document.add_paragraph('Exclude passwords, tokens, recovery links and real learner information. Keep unsynced device state for support.')
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Arial'; run.font.size = DocPt(10)
    document.save(HERE / 'RaPS-Beta-Feedback-Form.docx')


def build_logs():
    columns = ['report_id','reported_at_local','time_zone','test_installation','tester_alias','role',
               'device_model','os_version','browser_version','app_version','pwa_or_browser','network_state',
               'mission_id','synthetic_case_id','expected','actual','reproduction_steps','reproduced_count',
               'attempt_count','impact','independent_helped_blocked_not_run','data_state','screenshot_reference',
               'suggested_change','what_to_keep','owner','status','retest_version','retest_result']
    with (HERE / 'RaPS-Beta-Issue-Log.csv').open('w', newline='') as output:
        csv.writer(output).writerow(columns)
    with (HERE / 'RaPS-Beta-Readiness-Register.csv').open('w', newline='') as output:
        writer = csv.writer(output)
        writer.writerow(['check_id','check','owner_role','status','evidence','checked_at','follow_up'])
        for row in [
            ('R01','Matching source merged','Pilot owner','Verified in Git','Main merge '+SOURCE_COMMIT,'','Confirm release on devices separately'),
            ('R02','Web.8 NUC backend','NUC owner','Verified from supplied rollout log','Backend 8a9d138; legacy duplicate TEST catalog retained','','Web.9 does not migrate deployed records'),
            ('R03','Actual device/PWA version web.9','Site PM','Not yet verified','','','Check every participating device'),
            ('R04','Real username setup/reset/old-session scope','Enterprise owner','Not yet verified','','','Use a real pilot account on a second device'),
            ('R05','Known case two-device sync/final result/reports','SME + Site PM','Not yet verified','','','Allow time for existing requirements'),
            ('R06','Physical-device outage/timer recovery','Site PM','Pilot evidence required','','','Capture device and local/cloud state'),
            ('R07','Existing automatic closed-class retention','NUC owner','Not yet verified','','','Inspect an actual closed TEST class archive'),
            ('R08','Selected-class reporting CSV archive','NUC owner','Not yet verified','','','Verify selected TEST files and completion proof'),
            ('R09','Off-NUC backup and isolated restore','NUC owner','Not yet verified','','','Before meaningful record retention or expansion'),
            ('R10','Support route / nominees / dates','Pilot owner','To assign','','','Complete before briefing the cohort'),
            ('R11','Reported assessment-start failure','Pilot owner','Local fix; deployed retest pending','Shared regular/TEST INSERT RETURNING failure reproduced; protected insert/read fix','','Publish web.9 and retry the reported class before cohort launch')]:
            writer.writerow(row)


def main():
    register_pdf_fonts()
    slides = build_deck()
    build_markdown_pdf(HERE/'instructor-challenge-card.md', HERE/'RaPS-Instructor-Challenge-Card.pdf', 'Instructor challenge card')
    build_markdown_pdf(HERE/'facilitator-guide.md', HERE/'RaPS-Facilitator-Guide.pdf', 'Facilitator guide')
    build_feedback()
    build_logs()
    print(json.dumps({'slides': slides, 'presentation': PPTX_PATH.name,
                      'note': 'Convert presentation to PDF with LibreOffice; inspect outputs before distribution.'}))


if __name__ == '__main__':
    main()
